"""Document search API — search inside local Word files on disk."""
from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Query
from pydantic import BaseModel

from src.api._common import api_error, ok

router = APIRouter()


class SearchRequest(BaseModel):
    query: str
    root: str | None = None
    max_results: int = 50


class SearchResult(BaseModel):
    path: str
    filename: str
    snippet: str
    score: float


def _read_docx_text(path: Path) -> str:
    """Extract text from a .docx file."""
    try:
        from docx import Document  # noqa: E402

        doc = Document(path)
        parts: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        parts.append(text)
        return "\n".join(parts)
    except Exception:
        return ""


def _read_doc_text(path: Path) -> str:
    """Best-effort text extraction from legacy .doc files.

    This does not parse the full Word binary format; it extracts readable
    text blocks where possible. For full fidelity, prefer converting to
    .docx or using a dedicated parser.
    """
    try:
        data = path.read_bytes()
    except Exception:
        return ""
    text = ""
    current: list[str] = []
    i = 0
    while i < len(data):
        # crude printable run detection
        if 32 <= data[i] <= 126 or data[i] in (9, 10, 13):
            current.append(chr(data[i]))
            i += 1
        else:
            segment = "".join(current).strip()
            if len(segment) >= 4:
                text += segment + " "
            current = []
            i += 1
    segment = "".join(current).strip()
    if len(segment) >= 4:
        text += segment + " "
    return text.strip()


def _score_match(text: str, query: str) -> tuple[float, str]:
    """Simple relevance score + snippet extraction."""
    q = query.lower()
    positions = [i for i in range(len(text)) if text[i : i + len(query)].lower() == q]
    if not positions:
        return 0.0, ""
    score = float(len(positions)) / max(1, len(text))
    start = max(0, positions[0] - 80)
    end = min(len(text), positions[0] + len(query) + 120)
    snippet = text[start:end].replace("\n", " ")
    if start > 0:
        snippet = "..." + snippet
    if end < len(text):
        snippet = snippet + "..."
    return score, snippet


@router.post("/document-search")
def search_documents(req: SearchRequest) -> dict:
    query = req.query.strip()
    if not query:
        raise api_error("bad_request", "query is required", 400)

    root = Path(req.root or os.path.expanduser("~/Documents")).expanduser().resolve()
    if not root.exists() or not root.is_dir():
        raise api_error("bad_root", f"root folder not found: {root}", 400)

    max_depth = 6
    results: list[dict[str, Any]] = []
    for dirpath, _, filenames in os.walk(root):
        depth = len(Path(dirpath).resolve().parts) - len(root.parts)
        if depth > max_depth:
            continue
        for filename in filenames:
            lower_name = filename.lower()
            if not (lower_name.endswith(".docx") or lower_name.endswith(".doc")):
                continue
            full_path = Path(dirpath) / filename
            try:
                if lower_name.endswith(".docx"):
                    text = _read_docx_text(full_path)
                else:
                    text = _read_doc_text(full_path)
            except Exception:
                continue
            score, snippet = _score_match(text, query)
            if score <= 0:
                continue
            results.append(
                {
                    "path": str(full_path.resolve()),
                    "filename": filename,
                    "snippet": snippet,
                    "score": round(score, 4),
                }
            )
            if len(results) >= max(req.max_results, 200):
                break
        if len(results) >= max(req.max_results, 200):
            break

    results.sort(key=lambda item: item["score"], reverse=True)
    results = results[: max(req.max_results, 1)]
    return ok(
        {
            "query": query,
            "root": str(root),
            "results": results,
            "count": len(results),
        }
    )
